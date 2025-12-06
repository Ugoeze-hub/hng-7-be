import PyPDF2
from docx import Document as DocxDocument
import io
from fastapi import HTTPException, status

class TextExtractionService:
    @staticmethod
    def extract_from_pdf(file_content: bytes) -> str:
        """Extract text from PDF file."""
        try:
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_content))
            text = ""
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
            
            text = text.strip()
            
            if not text:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="No text could be extracted from the PDF"
                )
            
            return text
        except HTTPException:
            raise
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to extract text from PDF: {str(e)}"
            )

    @staticmethod
    def extract_from_docx(file_content: bytes) -> str:
        """Extract text from DOCX file."""
        try:
            doc = DocxDocument(io.BytesIO(file_content))
            text = []
            
            for paragraph in doc.paragraphs:
                if paragraph.text.strip(): 
                    text.append(paragraph.text)
            
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text.append(cell.text)
            
            full_text = "\n".join(text)
            
            if not full_text.strip():
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="No text could be extracted from the DOCX file"
                )
            
            return full_text
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to extract text from DOCX: {str(e)}"
            )
        

    @staticmethod
    def extract_text(file_content: bytes, filename: str) -> str:
        """Extract text from file based on extension."""
        filename_lower = filename.lower()
        
        if filename_lower.endswith('.pdf'):
            return TextExtractionService.extract_from_pdf(file_content)
        elif filename_lower.endswith('.docx'):
            return TextExtractionService.extract_from_docx(file_content)
        else:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Unsupported file format. Only PDF and DOCX are supported."
            )